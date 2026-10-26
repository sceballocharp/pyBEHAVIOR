"""Rebuild an Excel session log from session_history_*.json files.

Install once: python -m pip install openpyxl
Run: python update_animal_log.py
Optional: python update_animal_log.py --input-dir PATH --output PATH.xlsx
"""

import argparse
import json
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path


def read_sessions(folder):
    files = sorted(folder.glob('session_history_*.json'))
    if not files:
        raise ValueError(f'No session_history_*.json files found in {folder}')
    sessions = []
    seen = {}
    for path in files:
        entries = json.loads(path.read_text(encoding='utf-8-sig'))
        if not isinstance(entries, list):
            raise ValueError(f'{path.name}: expected a JSON list')
        for index, entry in enumerate(entries, 1):
            label = f'{path.name}, entry {index}'
            if not isinstance(entry, dict):
                raise ValueError(f'{label}: expected a session object')
            for key in ('started_at', 'animal_id', 'project'):
                if entry.get(key) in (None, ''):
                    raise ValueError(f'{label}: missing {key}')
            started = datetime.fromisoformat(entry['started_at'])
            if started.tzinfo is None:
                raise ValueError(f'{label}: started_at must include a timezone')
            session_id = entry.get('session_id')
            if session_id and session_id in seen:
                if seen[session_id] != entry:
                    raise ValueError(f'{label}: conflicting copies of session {session_id}')
                continue
            if session_id:
                seen[session_id] = entry
            weights = entry.get('weights') or {}
            if not isinstance(weights, dict):
                raise ValueError(f'{label}: weights must be an object')
            sequence = weights.get('sequence_weights')
            if isinstance(sequence, list):
                sequence = ' '.join(str(value) for value in sequence)
            row = [started.date(), str(entry['animal_id']), str(entry['project']),
                   sequence, entry.get('trials'), entry.get('rewards')]
            sessions.append((started.astimezone(timezone.utc), row))
    sessions.sort(key=lambda item: item[0])
    return [row for _, row in sessions], files


def write_log(rows, output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = ['Date', 'Animal', 'Protocol', 'Sequence weights', 'Trials', 'Rewards']
    animals = sorted({row[1] for row in rows})
    groups = [('Session log', rows)] + [
        (animal, [row for row in rows if row[1] == animal]) for animal in animals
    ]
    for number, (name, data) in enumerate(groups):
        # Excel sheet names cannot contain these characters or exceed 31 characters.
        title = re.sub(r'[\\/*?:\[\]]', '_', name).strip("'")[:31] or 'Animal'
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in data:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        for column, width in zip('ABCDEF', (16, 14, 28, 24, 12, 12)):
            sheet.column_dimensions[column].width = width
        for cells in sheet:
            sheet.row_dimensions[cells[0].row].height = 22
            for cell in cells:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='center')
                # Source strings must remain literal Excel values, even with '=' prefixes.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            if cells[0].row > 1:
                cells[0].number_format = 'dd/mm/yyyy'
                cells[1].number_format = '@'
                cells[3].number_format = '@'
                cells[4].number_format = cells[5].number_format = '0'
        for cell in sheet[1]:
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        if data:
            table = Table(displayName=f'SessionTable{number}', ref=f'A1:F{sheet.max_row}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)

    output.parent.mkdir(parents=True, exist_ok=True)
    # Save completely before replacing the previous report.
    handle, temporary = tempfile.mkstemp(suffix='.xlsx', dir=output.parent)
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return animals


def main():
    folder = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', type=Path, default=folder)
    parser.add_argument('--output', type=Path, default=folder / 'animal_log.xlsx')
    args = parser.parse_args()
    try:
        rows, files = read_sessions(args.input_dir)
        animals = write_log(rows, args.output)
    except ModuleNotFoundError as error:
        if error.name != 'openpyxl':
            raise
        print('Install the Excel dependency: python -m pip install openpyxl', file=sys.stderr)
        return 1
    except PermissionError:
        print(f'Cannot write {args.output}. Close it in Excel and check folder permissions, then retry.', file=sys.stderr)
        return 1
    except (OSError, ValueError, TypeError) as error:
        print(f'Could not update log: {error}', file=sys.stderr)
        return 1
    print(f'Updated {args.output}: {len(rows)} sessions, {len(animals)} animals, {len(files)} input files.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
